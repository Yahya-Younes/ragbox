"""Turn an uploaded file into text, page by page.

Every parser returns a list of (page_number | None, text). Files are checked by their
magic bytes, not only their extension, and .docx archives are checked against
decompression bombs before being opened.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

MAX_PDF_PAGES = 3000
MAX_DOCX_UNCOMPRESSED = 300 * 1024 * 1024  # 300 MB


class ParseError(ValueError):
    """The file can't be read: wrong type, damaged, encrypted or empty."""


def parse_file(path: Path, ext: str) -> list[tuple[int | None, str]]:
    ext = ext.lower()
    if ext == ".pdf":
        pages = _parse_pdf(path)
    elif ext == ".docx":
        pages = _parse_docx(path)
    elif ext in {".txt", ".md", ".markdown"}:
        pages = [(None, _read_text(path))]
    else:
        raise ParseError(f"Unsupported file type: {ext}")
    pages = [(p, t) for p, t in pages if t.strip()]
    if not pages:
        raise ParseError("No text found in this file (scanned PDFs need OCR, which ragbox doesn't do yet).")
    return pages


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    if b"\x00" in raw[:4096]:
        raise ParseError("This doesn't look like a text file.")
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError("Unknown text encoding.")  # pragma: no cover - latin-1 always decodes


def _parse_pdf(path: Path) -> list[tuple[int | None, str]]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    with path.open("rb") as f:
        if f.read(5) != b"%PDF-":
            raise ParseError("This file isn't a valid PDF.")
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ParseError("This PDF is password-protected.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ParseError(f"This PDF has more than {MAX_PDF_PAGES} pages.")
        return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
    except ParseError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise ParseError(f"This PDF couldn't be read ({exc.__class__.__name__}).") from exc


def _parse_docx(path: Path) -> list[tuple[int | None, str]]:
    if not zipfile.is_zipfile(path):
        raise ParseError("This file isn't a valid Word document.")
    with zipfile.ZipFile(path) as z:
        if sum(i.file_size for i in z.infolist()) > MAX_DOCX_UNCOMPRESSED:
            raise ParseError("This Word document is too large once uncompressed.")
        if "word/document.xml" not in z.namelist():
            raise ParseError("This file isn't a valid Word document.")

    import docx

    document = docx.Document(str(path))
    parts: list[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(c for c in cells if c))
    return [(None, "\n".join(parts))]
