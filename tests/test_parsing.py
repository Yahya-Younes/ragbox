import zipfile

import pytest

from ragbox.parsing import ParseError, parse_file


def make_pdf(text: str) -> bytes:
    """Build a minimal valid one-page PDF (with a correct xref table) containing `text`."""
    stream = f"BT /F1 18 Tf 20 100 Td ({text}) Tj ET".encode()
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 144]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length %d>>stream\n" % len(stream) + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj" % i + obj + b"endobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF" % (len(objects) + 1, xref)
    return out


def test_pdf(tmp_path):
    f = tmp_path / "a.pdf"
    f.write_bytes(make_pdf("Hello ragbox"))
    pages = parse_file(f, ".pdf")
    assert pages[0][0] == 1 and "Hello ragbox" in pages[0][1]


def test_pdf_with_wrong_content_is_rejected(tmp_path):
    f = tmp_path / "fake.pdf"
    f.write_text("not a pdf")
    with pytest.raises(ParseError):
        parse_file(f, ".pdf")


def test_docx_paragraphs_and_tables(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("Quarterly report")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Revenue", "42"
    path = tmp_path / "r.docx"
    d.save(path)
    ((_, text),) = parse_file(path, ".docx")
    assert "Quarterly report" in text and "Revenue | 42" in text


def test_docx_that_is_not_a_zip_is_rejected(tmp_path):
    f = tmp_path / "x.docx"
    f.write_text("hello")
    with pytest.raises(ParseError):
        parse_file(f, ".docx")


def test_zip_without_word_document_is_rejected(tmp_path):
    f = tmp_path / "x.docx"
    with zipfile.ZipFile(f, "w") as z:
        z.writestr("other.txt", "hi")
    with pytest.raises(ParseError):
        parse_file(f, ".docx")


def test_text_encodings(tmp_path):
    f = tmp_path / "n.txt"
    f.write_bytes("Café crème".encode("cp1252"))
    assert parse_file(f, ".txt")[0][1] == "Café crème"


def test_binary_and_empty_files_are_rejected(tmp_path):
    b = tmp_path / "b.txt"
    b.write_bytes(b"\x00\x01\x02")
    e = tmp_path / "e.md"
    e.write_text("   \n")
    for path in (b, e):
        with pytest.raises(ParseError):
            parse_file(path, path.suffix)
